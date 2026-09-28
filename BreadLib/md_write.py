from .md_notif import *
import pandas as pd
import os
from openpyxl.styles import Alignment


def write_daily_sheets(workbook, data, kw_number, column_names):
    days = list(dict.fromkeys(key[0] for key in data.keys()))
    sheets = {}
    rows_cursor = {}

    for index, day in enumerate(days):
        sheet = workbook.create_sheet(title=str(day), index=index)
        sheets[day] = sheet
        rows_cursor[day] = 3
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
        sheet.cell(row=1, column=1).value = f"KW{kw_number} - {day}"
        sheet.cell(row=1, column=1).alignment = Alignment(horizontal='center')
        for col_num, header in enumerate(column_names, 1):
            sheet.cell(row=2, column=col_num).value = header

    for (d, lv, mat), val in data.items():
        sheet = sheets[d]
        r = rows_cursor[d]
        sheet.cell(row=r, column=1).value = lv
        sheet.cell(row=r, column=2).value = mat
        sheet.cell(row=r, column=3).value = val.get('Count', 0)
        sheet.cell(row=r, column=4).value = val.get('Unit', 'm')
        rows_cursor[d] += 1


def write_summary(workbook, data_total, kw_number, column_names):
    sheet = workbook.create_sheet(title="Summary")
    sheet.cell(row=1, column=1).value = f"KW{kw_number} Summary"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)

    for col, text in enumerate(column_names, 1):
        sheet.cell(row=2, column=col).value = text

    for row, ((lv, mat), val) in enumerate(data_total.items(), start=3):
        sheet.cell(row=row, column=1).value = lv
        sheet.cell(row=row, column=2).value = mat.strip()
        sheet.cell(row=row, column=3).value = val.get('Count', 0)
        sheet.cell(row=row, column=4).value = val.get('Unit', 'm')


def build_table(path, data, kw_number, data_total=None):
    try:
        if not data:
            msg_error("No data provided to output!")
            return

        excel_file = os.path.join(path, "output.xlsx")
        column_names = ['LV-Position', 'Bezeichnung', 'Menge', 'Einheit']

        with pd.ExcelWriter(excel_file, engine="openpyxl") as writer:
            write_daily_sheets(writer.book, data, kw_number, column_names)

            if data_total is not None:
                write_summary(writer.book, data_total, kw_number, column_names)

            if "Sheet1" in writer.book.sheetnames:
                writer.book.remove(writer.book["Sheet1"])

        msg_info(f"Data saved with path: {excel_file}")
    except PermissionError:
        msg_warning("Close output file before processing!")
    except Exception as e:
        msg_warning(e)